"""Data format conversion: uploaded formats must become the engine CSV columns."""
from __future__ import annotations

import csv
import json
import tempfile
import unittest
from pathlib import Path

from workers.capabilities import dataset_import
from workers.local_compute import normalize_dataset

HEADERS = ["内容", "发布时间", "渠道", "标签"]
TEXTS = ["包裹配送延迟 物流 快递 延误", "退款 账单 重复扣费 客服", "登录 密码 账号 闪退"]
ROWS = 24


def read_rows(path: Path) -> tuple[list[str], list[dict[str, str]]]:
    with path.open(encoding="utf-8", newline="") as handle:
        reader = csv.DictReader(handle)
        return list(reader.fieldnames or []), list(reader)


class DatasetConversionTest(unittest.TestCase):
    def setUp(self) -> None:
        self.root = Path(tempfile.mkdtemp(prefix="theta-conversion-"))
        self.uploads = self.root / "uploads"

    def write_xlsx(self) -> Path:
        import openpyxl

        path = self.root / "反馈数据.xlsx"
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "反馈"
        sheet.append(HEADERS)
        for index in range(ROWS):
            sheet.append([TEXTS[index % len(TEXTS)], "2026-0%d-15" % (index % 3 + 1), "phone" if index % 2 else "web", "L%d" % (index % 2)])
        workbook.save(path)
        return path

    def write_jsonl(self) -> Path:
        path = self.root / "feedback.jsonl"
        records = [
            json.dumps({"body": TEXTS[index % len(TEXTS)], "created_at": "2026-0%d-15" % (index % 3 + 1), "channel": "phone" if index % 2 else "web"}, ensure_ascii=False)
            for index in range(ROWS)
        ]
        path.write_text("\n".join(records) + "\n", encoding="utf-8")
        return path

    def imported(self, source: Path) -> dict:
        return dataset_import({"filePath": str(source), "uploadDir": str(self.uploads)})

    def test_excel_with_chinese_headers_maps_every_engine_column(self) -> None:
        dataset = self.imported(self.write_xlsx())
        plan = {"modelId": "dtm", "textColumn": "内容", "timeColumn": "发布时间", "labelColumn": "标签", "covariates": ["渠道"], "params": {}}
        destination = self.root / "normalized.csv"
        normalize_dataset({"dataset": dataset, "plan": plan}, destination)
        header, rows = read_rows(destination)
        self.assertEqual(header, ["text", "label", "year", "timestamp", "cov_0"])
        self.assertEqual(len(rows), ROWS)
        self.assertEqual(rows[0]["text"], TEXTS[0])
        self.assertEqual(rows[0]["cov_0"], "web")
        self.assertEqual(rows[0]["year"], rows[0]["timestamp"])
        self.assertEqual(rows[1]["label"], "L1")

    def test_jsonl_records_map_text_and_covariates(self) -> None:
        dataset = self.imported(self.write_jsonl())
        plan = {"modelId": "stm", "textColumn": "body", "covariates": ["channel"], "params": {}}
        destination = self.root / "normalized.csv"
        normalize_dataset({"dataset": dataset, "plan": plan}, destination)
        header, rows = read_rows(destination)
        self.assertEqual(header, ["text", "cov_0"])
        self.assertEqual(len(rows), ROWS)
        self.assertTrue(all(row["cov_0"] in {"web", "phone"} for row in rows))
        self.assertTrue(all(row["text"] for row in rows))

    def test_table_formats_preserve_explicit_time_and_covariates_for_stm_and_dtm(self):
        import pyarrow as pa
        import pyarrow.parquet as pq
        records = [{'正文': TEXTS[index % 3], '发布时间': f'2026-0{index % 3 + 1}-15', '公众号名称': f'渠道{index % 2}'} for index in range(ROWS)]
        for suffix in ['csv', 'tsv', 'json', 'jsonl', 'ndjson', 'parquet']:
            source = self.root / ('多语言表格.' + suffix)
            if suffix == 'parquet':
                pq.write_table(pa.Table.from_pylist(records), source)
            elif suffix == 'json':
                source.write_text(json.dumps(records, ensure_ascii=False), encoding='utf-8')
            elif suffix in {'jsonl', 'ndjson'}:
                source.write_text('\n'.join(json.dumps(row, ensure_ascii=False) for row in records), encoding='utf-8')
            else:
                with source.open('w', encoding='utf-8-sig', newline='') as handle:
                    writer = csv.DictWriter(handle, fieldnames=list(records[0]), delimiter='\t' if suffix == 'tsv' else ',')
                    writer.writeheader(); writer.writerows(records)
            for model in ['stm', 'dtm']:
                with self.subTest(suffix=suffix, model=model):
                    destination = self.root / f'{model}-{suffix}.csv'
                    normalize_dataset({'dataset': self.imported(source), 'plan': {'modelId': model, 'textColumn': '正文', 'timeColumn': '发布时间', 'covariates': ['公众号名称'], 'params': {}}}, destination)
                    header, rows = read_rows(destination)
                    self.assertEqual(header, ['text', 'year', 'timestamp', 'cov_0'])
                    self.assertEqual(len(rows), ROWS)
                    self.assertEqual(rows[0]['text'], records[0]['正文'])
                    self.assertEqual(rows[0]['timestamp'], records[0]['发布时间'])
                    self.assertEqual(rows[0]['cov_0'], records[0]['公众号名称'])

    def test_original_column_names_never_reach_the_engine(self) -> None:
        dataset = self.imported(self.write_xlsx())
        destination = self.root / "normalized.csv"
        normalize_dataset({"dataset": dataset, "plan": {"modelId": "lda", "textColumn": "内容", "params": {}}}, destination)
        header, rows = read_rows(destination)
        self.assertEqual(header, ["text"])
        self.assertEqual(len(rows), ROWS)

    def test_a_changed_managed_copy_is_rejected(self) -> None:
        dataset = self.imported(self.write_jsonl())
        Path(dataset["managedPath"]).write_text("text\nchanged\n", encoding="utf-8")
        with self.assertRaises(ValueError):
            normalize_dataset({"dataset": dataset, "plan": {"modelId": "lda", "textColumn": "body", "params": {}}}, self.root / "out.csv")

    def test_gbk_csv_with_chinese_headers_is_decoded_and_mapped(self) -> None:
        path = self.root / "反馈-gbk.csv"
        path.write_bytes("内容,渠道\n包裹配送延迟 物流,web\n退款 重复扣费,phone\n".encode("gbk"))
        dataset = self.imported(path)
        destination = self.root / "normalized.csv"
        normalize_dataset({"dataset": dataset, "plan": {"modelId": "lda", "textColumn": "内容", "covariates": ["渠道"], "params": {}}}, destination)
        header, rows = read_rows(destination)
        self.assertEqual(header, ["text", "cov_0"])
        self.assertEqual(rows[0]["text"], "包裹配送延迟 物流")
        self.assertEqual(rows[1]["cov_0"], "phone")

    def test_utf8_bom_is_not_part_of_the_column_name(self) -> None:
        path = self.root / "bom.csv"
        path.write_bytes("\ufeff内容\n包裹配送延迟 物流\n".encode("utf-8"))
        dataset = self.imported(path)
        destination = self.root / "normalized.csv"
        normalize_dataset({"dataset": dataset, "plan": {"modelId": "lda", "textColumn": "内容", "params": {}}}, destination)
        header, rows = read_rows(destination)
        self.assertEqual(header, ["text"])
        self.assertEqual(rows[0]["text"], "包裹配送延迟 物流")

    def test_tab_separated_records_are_split_into_columns(self) -> None:
        path = self.root / "feedback.tsv"
        path.write_text("内容\t渠道\n包裹配送延迟\tweb\n", encoding="utf-8")
        dataset = self.imported(path)
        destination = self.root / "normalized.csv"
        normalize_dataset({"dataset": dataset, "plan": {"modelId": "lda", "textColumn": "内容", "covariates": ["渠道"], "params": {}}}, destination)
        header, rows = read_rows(destination)
        self.assertEqual(header, ["text", "cov_0"])
        self.assertEqual(rows[0]["cov_0"], "web")

    def test_docx_paragraphs_are_available_as_training_rows(self) -> None:
        try:
            from docx import Document
        except ImportError:
            self.skipTest('python-docx is not installed in this worker profile')
        source = self.root / "feedback.docx"
        document = Document()
        document.add_paragraph("包裹配送延迟 物流 快递 延误")
        document.add_paragraph("退款 账单 重复扣费 客服")
        document.save(source)
        dataset = self.imported(source)
        destination = self.root / "normalized.csv"
        normalize_dataset({"dataset": dataset, "plan": {"modelId": "lda", "textColumn": "text", "params": {}}}, destination)
        header, rows = read_rows(destination)
        self.assertEqual(header, ["text"])
        self.assertEqual([row["text"] for row in rows], ["包裹配送延迟 物流 快递 延误", "退款 账单 重复扣费 客服"])


if __name__ == "__main__":
    unittest.main()
