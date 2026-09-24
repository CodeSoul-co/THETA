import tempfile
import unittest
from pathlib import Path
from workers.capabilities import dataset_import, dataset_preview

class DatasetPreviewTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)

    def preview(self, file):
        return dataset_preview({'dataset': dataset_import({'filePath': str(file), 'uploadDir': str(self.root / 'uploads')})})

    def test_text_preview_reads_first_records_without_column_selection(self):
        for suffix in ['.txt', '.md']:
            file = self.root / ('正文' + suffix)
            file.write_text('\n'.join('正文第%d行' % n for n in range(100)), encoding='utf8')
            preview = self.preview(file)
            self.assertEqual(preview['inputKind'], 'text')
            self.assertEqual(preview['textColumn'], 'text')
            self.assertEqual(preview['totalRecords'], 100)
            self.assertEqual([x['text'] for x in preview['segments']], ['正文第%d行' % n for n in range(5)])

    def test_structured_data_retains_columns_and_zero_values(self):
        file = self.root / 'table.csv'
        file.write_text('text,value\n正文,0\n', encoding='utf8')
        preview = self.preview(file)
        self.assertEqual(preview['inputKind'], 'table')
        self.assertEqual(preview['rows'], [['正文', '0']])
        self.assertNotIn('segments', preview)

    def test_docx_preview_uses_paragraphs(self):
        from docx import Document
        file = self.root / 'paragraphs.docx'
        doc = Document(); doc.add_paragraph('第一段'); doc.add_paragraph(''); doc.add_paragraph('第三段'); doc.save(file)
        preview = self.preview(file)
        self.assertEqual(preview['segments'], [{'text': '第一段', 'paragraph': 1}, {'text': '第三段', 'paragraph': 3}])

    def test_empty_pdf_explains_text_recognition(self):
        from PyPDF2 import PdfWriter
        file = self.root / 'scan.pdf'
        writer = PdfWriter(); writer.add_blank_page(width=100, height=100)
        with file.open('wb') as handle: writer.write(handle)
        with self.assertRaisesRegex(ValueError, '文字识别'):
            self.preview(file)

    def test_office_files_in_unicode_paths_and_invalid_containers(self):
        import zipfile
        from docx import Document
        from openpyxl import Workbook
        folder = self.root / '中文用户' / '安装目录 数据'
        folder.mkdir(parents=True)
        word = folder / '用户原文.docx'
        doc = Document(); doc.add_paragraph('真实正文'); doc.save(word)
        excel = folder / '上传表格.xlsx'
        workbook = Workbook(); workbook.active.append(['正文']); workbook.active.append(['真实正文']); workbook.save(excel)
        for file in [word, excel]:
            with zipfile.ZipFile(file, 'a') as archive:
                archive.writestr('padding.bin', b'x' * (12 * 1024 * 1024), compress_type=zipfile.ZIP_STORED)
            self.assertEqual(self.preview(file)['rows'][0][0], '真实正文')
        for suffix in ['.docx', '.xlsx']:
            file = folder / ('broken' + suffix)
            for content, message in [(b'PK\x03\x04truncated', '不完整'), (b'plain text', '扩展名'), (b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1old office', '加密')]:
                file.write_bytes(content)
                with self.assertRaisesRegex(ValueError, message): self.preview(file)
            with zipfile.ZipFile(file, 'w') as archive: archive.writestr('wrong.xml', 'not an Office document')
            with self.assertRaisesRegex(ValueError, '不匹配'): self.preview(file)

    def test_word_tables_and_long_paragraphs_keep_real_source_positions(self):
        from docx import Document
        from workers.dataset.readers import load_dataset
        file = self.root / '文档分段.docx'
        doc = Document(); doc.add_paragraph('标题')
        text = '这是一段真实语句，按句子边界切分。' * 180
        doc.add_paragraph(text)
        table = doc.add_table(rows=1, cols=2); table.cell(0, 0).text = '指标'; table.cell(0, 1).text = '正文内容'
        doc.save(file)
        rows = load_dataset(file).rows
        chunks = [row for row in rows if row.get('paragraph') == 2]
        self.assertGreater(len(chunks), 1)
        self.assertEqual(''.join(row['text'] for row in chunks), text)
        self.assertTrue(all(len(row['text']) <= 1600 for row in chunks))
        self.assertEqual(rows[-1]['table'], 1)
        self.assertIn('正文内容', rows[-1]['text'])

    def test_folder_collection_retains_all_documents_and_source_names(self):
        from workers.dataset.collection import combine
        files = []
        for index in range(3):
            source = self.root / f'doc-{index}.txt'
            source.write_text(f'正文 {index}\n第二段 {index}', encoding='utf8')
            files.append(dataset_import({'filePath': str(source), 'uploadDir': str(self.root / 'uploads')}))
        result = combine({'datasets': files, 'uploadDir': str(self.root / 'uploads')})
        self.assertEqual(result['recordCount'], 6)
        preview = dataset_preview({'dataset': result})
        self.assertEqual(preview['inputKind'], 'text')
        self.assertEqual(preview['totalRecords'], 6)
        self.assertEqual(preview['segments'][0]['source_file'], 'doc-0.txt')
        self.assertEqual(preview['segments'][2]['source_file'], 'doc-1.txt')
        source = self.root / 'table.csv'; source.write_text('text\n正文', encoding='utf8')
        table = dataset_import({'filePath': str(source), 'uploadDir': str(self.root / 'uploads')})
        with self.assertRaisesRegex(ValueError, '表格请单独上传'):
            combine({'datasets': [table], 'uploadDir': str(self.root / 'uploads')})
