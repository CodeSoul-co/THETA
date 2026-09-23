from __future__ import annotations

import csv
import json
import tempfile
import unittest
from pathlib import Path

from worker.dataset import normalize_dataset, split_parameters
from worker.errors import PermanentJobError


class DatasetNormalizationTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name)

    def tearDown(self) -> None:
        self.temporary.cleanup()

    def normalized(self, source: Path, file_format: str, options: dict) -> list[dict[str, str]]:
        output = self.root / "normalized.csv"
        normalize_dataset(source, output, file_format, options)
        with output.open(encoding="utf-8", newline="") as handle:
            return list(csv.DictReader(handle))

    def test_jsonl_maps_all_engine_columns(self) -> None:
        source = self.root / "records.jsonl"
        source.write_text(
            "\n".join(
                json.dumps({"内容": f"文本 {index}", "时间": 2020 + index, "标签": "A", "渠道": "web"}, ensure_ascii=False)
                for index in range(2)
            ),
            encoding="utf-8",
        )
        rows = self.normalized(source, "jsonl", {
            "text_column": "内容", "time_column": "时间", "label_column": "标签", "covariates": ["渠道"]
        })
        self.assertEqual(list(rows[0]), ["text", "label", "year", "timestamp", "cov_0"])
        self.assertEqual(rows[1]["text"], "文本 1")
        self.assertEqual(rows[0]["year"], rows[0]["timestamp"])

    def test_excel_sheet_and_parquet_are_supported_when_adapters_are_installed(self) -> None:
        try:
            import pandas as pd
            import openpyxl  # noqa: F401
            import pyarrow  # noqa: F401
        except ImportError:
            self.skipTest("optional tabular adapters are not installed")
        frame = pd.DataFrame({"body": ["first", "second"], "group": ["x", "y"]})
        workbook = self.root / "records.xlsx"
        with pd.ExcelWriter(workbook) as writer:
            frame.to_excel(writer, sheet_name="data", index=False)
        parquet = self.root / "records.parquet"
        frame.to_parquet(parquet, index=False)
        for source, file_format in ((workbook, "xlsx"), (parquet, "parquet")):
            with self.subTest(file_format=file_format):
                rows = self.normalized(source, file_format, {"text_column": "body", "covariates": ["group"], "sheet": "data"})
                self.assertEqual(rows[0], {"text": "first", "cov_0": "x"})

    def test_text_lines_and_gbk_csv_are_supported(self) -> None:
        text = self.root / "records.txt"
        text.write_text("第一条\n第二条\n", encoding="utf-8")
        self.assertEqual([row["text"] for row in self.normalized(text, "txt", {})], ["第一条", "第二条"])
        csv_file = self.root / "records.csv"
        csv_file.write_bytes("内容;渠道\n退款;电话\n".encode("gbk"))
        self.assertEqual(
            self.normalized(csv_file, "csv", {"text_column": "内容", "covariates": ["渠道"]})[0],
            {"text": "退款", "cov_0": "电话"},
        )

    def test_utf8_probe_ending_inside_a_chinese_character_is_not_misdetected_as_latin1(self) -> None:
        source = self.root / "large.csv"
        # Header plus ASCII payload place the first byte of 中 exactly at the
        # end of the 65,536-byte encoding probe.
        text = "a" * 65_530 + "中文语料"
        source.write_bytes(("text\n" + text + "\n").encode("utf-8"))
        rows = self.normalized(source, "csv", {})
        self.assertTrue(rows[0]["text"].endswith("中文语料"))

    def test_normalized_csv_is_idempotent(self) -> None:
        source = self.root / "data.csv"
        source.write_text(
            "text,label,year,timestamp,cov_1,cov_0\n退款,A,2024,2024,华东,电话\n",
            encoding="utf-8",
        )
        rows = self.normalized(source, "csv", {})
        self.assertEqual(
            rows[0],
            {"text": "退款", "label": "A", "year": "2024", "timestamp": "2024", "cov_0": "电话", "cov_1": "华东"},
        )

    def test_explicit_empty_mappings_override_normalized_column_preservation(self) -> None:
        source = self.root / "data.csv"
        source.write_text("text,label,timestamp,cov_0\n退款,A,2024,电话\n", encoding="utf-8")
        rows = self.normalized(source, "csv", {"label_column": None, "time_column": None, "covariates": []})
        self.assertEqual(rows[0], {"text": "退款"})

    def test_markdown_and_docx_are_converted_to_text_records(self) -> None:
        markdown = self.root / "notes.md"
        markdown.write_text("第一段\n\n第二段\n", encoding="utf-8")
        self.assertEqual([row["text"] for row in self.normalized(markdown, "md", {})], ["第一段", "第二段"])
        try:
            from docx import Document
        except ImportError:
            self.skipTest("python-docx is not installed")
        document = Document()
        document.add_paragraph("退款处理说明")
        document.add_paragraph("物流延迟说明")
        path = self.root / "notes.docx"
        document.save(path)
        self.assertEqual(
            [row["text"] for row in self.normalized(path, "docx", {})],
            ["退款处理说明", "物流延迟说明"],
        )

    def test_unknown_columns_and_formats_fail_before_training(self) -> None:
        source = self.root / "records.json"
        source.write_text('[{"body":"text"}]', encoding="utf-8")
        with self.assertRaisesRegex(PermanentJobError, "text column"):
            self.normalized(source, "json", {"text_column": "missing"})
        with self.assertRaisesRegex(PermanentJobError, "unsupported dataset format"):
            self.normalized(source, "sqlite", {})

    def test_input_options_never_reach_engine_parameters(self) -> None:
        options, params = split_parameters({
            "input.text_column": "body", "input.covariates": ["group"], "num_topics": 8, "model.alpha": 0.2
        })
        self.assertEqual(options, {"text_column": "body", "covariates": ["group"]})
        self.assertEqual(params, {"num_topics": 8, "model.alpha": 0.2})


if __name__ == "__main__":
    unittest.main()
